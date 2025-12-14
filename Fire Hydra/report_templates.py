"""
Report Templates Module - Custom Report Generation
FireHydra - Yangın Söndürme Sistemi Hidrolik Hesaplama
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser
from dataclasses import dataclass, field
from typing import Optional, Dict, List
from datetime import datetime
import os

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm, cm
    from reportlab.pdfgen import canvas
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from jinja2 import Template
    HAS_JINJA2 = True
except ImportError:
    HAS_JINJA2 = False


@dataclass
class ReportTemplate:
    """Rapor şablonu"""
    name: str
    description: str
    format: str  # 'pdf', 'docx', 'html'
    
    # Görsel öğeler
    include_logo: bool = True
    logo_path: str = ""
    include_cover_page: bool = True
    include_toc: bool = True
    include_appendix: bool = True
    
    # İçerik bölümleri
    include_summary: bool = True
    include_network_diagram: bool = True
    include_calculation_details: bool = True
    include_pipe_table: bool = True
    include_node_table: bool = True
    include_validation_report: bool = True
    include_material_list: bool = True
    include_cost_estimate: bool = True
    
    # Şirket bilgileri
    company_name: str = ""
    company_address: str = ""
    company_phone: str = ""
    company_email: str = ""
    company_web: str = ""
    
    # Proje bilgileri
    project_name: str = ""
    project_number: str = ""
    project_location: str = ""
    client_name: str = ""
    engineer_name: str = ""
    reviewer_name: str = ""
    
    # Stil ayarları
    header_color: str = "#1f4788"
    accent_color: str = "#4472c4"
    font_family: str = "Helvetica"
    font_size: int = 10


@dataclass
class ReportSection:
    """Rapor bölümü"""
    title: str
    content: str
    include: bool = True
    page_break: bool = False


class PDFReportGenerator:
    """PDF rapor oluşturucu"""
    
    def __init__(self, template: ReportTemplate, network, results, validation_results=None):
        self.template = template
        self.network = network
        self.results = results
        self.validation_results = validation_results
        
        if not HAS_REPORTLAB:
            raise ImportError("reportlab kütüphanesi gerekli. 'pip install reportlab' çalıştırın.")
    
    def generate(self, filepath: str) -> bool:
        """Rapor oluştur"""
        try:
            doc = SimpleDocTemplate(filepath, pagesize=A4,
                                   leftMargin=2*cm, rightMargin=2*cm,
                                   topMargin=2*cm, bottomMargin=2*cm)
            
            story = []
            styles = getSampleStyleSheet()
            
            # Özel stil oluştur
            self._create_custom_styles(styles)
            
            # Kapak sayfası
            if self.template.include_cover_page:
                story.extend(self._create_cover_page(styles))
                story.append(PageBreak())
            
            # İçindekiler (placeholder)
            if self.template.include_toc:
                story.extend(self._create_toc(styles))
                story.append(PageBreak())
            
            # Özet
            if self.template.include_summary:
                story.extend(self._create_summary(styles))
                story.append(Spacer(1, 1*cm))
            
            # Network diyagramı
            if self.template.include_network_diagram:
                story.extend(self._create_network_diagram(styles))
                story.append(Spacer(1, 1*cm))
            
            # Hesaplama detayları
            if self.template.include_calculation_details:
                story.extend(self._create_calculation_details(styles))
                story.append(Spacer(1, 1*cm))
            
            # Boru tablosu
            if self.template.include_pipe_table:
                story.extend(self._create_pipe_table(styles))
                story.append(Spacer(1, 1*cm))
            
            # Node tablosu
            if self.template.include_node_table:
                story.extend(self._create_node_table(styles))
                story.append(Spacer(1, 1*cm))
            
            # Doğrulama raporu
            if self.template.include_validation_report and self.validation_results:
                story.extend(self._create_validation_report(styles))
                story.append(Spacer(1, 1*cm))
            
            # Malzeme listesi
            if self.template.include_material_list:
                story.extend(self._create_material_list(styles))
                story.append(Spacer(1, 1*cm))
            
            # PDF oluştur
            doc.build(story)
            return True
            
        except Exception as e:
            print(f"PDF oluşturma hatası: {e}")
            return False
    
    def _create_custom_styles(self, styles):
        """Özel stiller oluştur"""
        # Başlık stili
        styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor(self.template.header_color),
            spaceAfter=30,
            alignment=TA_CENTER
        ))
        
        # Alt başlık stili
        styles.add(ParagraphStyle(
            name='CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor(self.template.accent_color),
            spaceAfter=12
        ))
    
    def _create_cover_page(self, styles):
        """Kapak sayfası"""
        elements = []
        
        # Logo
        if self.template.include_logo and self.template.logo_path and os.path.exists(self.template.logo_path):
            try:
                logo = Image(self.template.logo_path, width=4*cm, height=4*cm)
                elements.append(logo)
                elements.append(Spacer(1, 2*cm))
            except:
                pass
        
        # Başlık
        title = Paragraph(f"<b>{self.template.project_name}</b>", styles['CustomTitle'])
        elements.append(title)
        elements.append(Spacer(1, 1*cm))
        
        # Alt başlık
        subtitle = Paragraph("Yangın Söndürme Sistemi<br/>Hidrolik Hesap Raporu", styles['CustomHeading'])
        elements.append(subtitle)
        elements.append(Spacer(1, 3*cm))
        
        # Proje bilgileri tablosu
        data = []
        if self.template.project_number:
            data.append(['Proje No:', self.template.project_number])
        if self.template.project_location:
            data.append(['Lokasyon:', self.template.project_location])
        if self.template.client_name:
            data.append(['Müşteri:', self.template.client_name])
        if self.template.engineer_name:
            data.append(['Hazırlayan:', self.template.engineer_name])
        if self.template.reviewer_name:
            data.append(['Kontrol Eden:', self.template.reviewer_name])
        data.append(['Tarih:', datetime.now().strftime('%d.%m.%Y')])
        
        if data:
            t = Table(data, colWidths=[5*cm, 10*cm])
            t.setStyle(TableStyle([
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ]))
            elements.append(t)
        
        # Şirket bilgileri (alt kısım)
        elements.append(Spacer(1, 3*cm))
        if self.template.company_name:
            company_info = []
            company_info.append(f"<b>{self.template.company_name}</b>")
            if self.template.company_address:
                company_info.append(self.template.company_address)
            if self.template.company_phone:
                company_info.append(f"Tel: {self.template.company_phone}")
            if self.template.company_email:
                company_info.append(f"E-posta: {self.template.company_email}")
            if self.template.company_web:
                company_info.append(self.template.company_web)
            
            company_text = Paragraph("<br/>".join(company_info), styles['Normal'])
            elements.append(company_text)
        
        return elements
    
    def _create_toc(self, styles):
        """İçindekiler"""
        elements = []
        elements.append(Paragraph("<b>İÇİNDEKİLER</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        toc_items = []
        page_num = 3  # Başlangıç sayfa numarası
        
        if self.template.include_summary:
            toc_items.append(f"1. Özet .......................................................... {page_num}")
            page_num += 1
        if self.template.include_calculation_details:
            toc_items.append(f"2. Hesaplama Detayları ........................................ {page_num}")
            page_num += 1
        if self.template.include_pipe_table:
            toc_items.append(f"3. Boru Tablosu ............................................... {page_num}")
            page_num += 1
        if self.template.include_node_table:
            toc_items.append(f"4. Node Tablosu ............................................... {page_num}")
            page_num += 1
        if self.template.include_validation_report:
            toc_items.append(f"5. Doğrulama Raporu ........................................... {page_num}")
            page_num += 1
        if self.template.include_material_list:
            toc_items.append(f"6. Malzeme Listesi ............................................ {page_num}")
            page_num += 1
        
        for item in toc_items:
            elements.append(Paragraph(item, styles['Normal']))
            elements.append(Spacer(1, 0.3*cm))
        
        return elements
    
    def _create_summary(self, styles):
        """Özet bölümü"""
        elements = []
        elements.append(Paragraph("<b>1. ÖZET</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        summary_data = [
            ['Parametre', 'Değer'],
            ['Toplam Debi', f"{self.results.total_flow:.1f} L/dk" if hasattr(self.results, 'total_flow') else 'N/A'],
            ['Toplam Basınç', f"{self.results.total_pressure:.2f} Bar" if hasattr(self.results, 'total_pressure') else 'N/A'],
            ['Maksimum Hız', f"{self.results.max_velocity:.2f} m/s" if hasattr(self.results, 'max_velocity') else 'N/A'],
            ['Toplam Boru Sayısı', str(len(self.network.pipes))],
            ['Toplam Node Sayısı', str(len(self.network.nodes))],
            ['Sprinkler Sayısı', str(len([n for n in self.network.nodes if n.node_type == 'sprinkler']))],
        ]
        
        t = Table(summary_data, colWidths=[8*cm, 8*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(self.template.header_color)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ]))
        elements.append(t)
        
        return elements
    
    def _create_network_diagram(self, styles):
        """Network diyagramı (placeholder)"""
        elements = []
        elements.append(Paragraph("<b>Network Diyagramı</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        elements.append(Paragraph("(Network görselleştirmesi buraya eklenecek)", styles['Normal']))
        return elements
    
    def _create_calculation_details(self, styles):
        """Hesaplama detayları"""
        elements = []
        elements.append(Paragraph("<b>2. HESAPLAMA DETAYLARI</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        # Hesaplama parametreleri
        params_text = f"""
        <b>Hesaplama Yöntemi:</b> Hazen-Williams<br/>
        <b>Hazen-Williams Katsayısı (C):</b> 120<br/>
        <b>Tasarım Yoğunluğu:</b> {getattr(self.results, 'density', 'N/A')} L/dk/m²<br/>
        <b>İşletme Alanı:</b> {getattr(self.results, 'operating_area', 'N/A')} m²<br/>
        <b>Hesaplama Tarihi:</b> {datetime.now().strftime('%d.%m.%Y %H:%M')}
        """
        elements.append(Paragraph(params_text, styles['Normal']))
        
        return elements
    
    def _create_pipe_table(self, styles):
        """Boru tablosu"""
        elements = []
        elements.append(Paragraph("<b>3. BORU TABLOSU</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        # Tablo başlıkları
        data = [['Boru ID', 'Başlangıç', 'Bitiş', 'Çap (mm)', 'Uzunluk (m)', 'Debi (L/dk)', 'Hız (m/s)']]
        
        # Boru verileri
        for pipe in self.network.pipes:
            row = [
                pipe.id,
                pipe.start_node_id,
                pipe.end_node_id,
                f"{pipe.diameter_mm:.0f}",
                f"{pipe.length:.2f}",
                f"{getattr(pipe, 'flow', 0):.1f}",
                f"{getattr(pipe, 'velocity', 0):.2f}"
            ]
            data.append(row)
        
        t = Table(data, colWidths=[2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2.5*cm, 2*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(self.template.header_color)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        elements.append(t)
        
        return elements
    
    def _create_node_table(self, styles):
        """Node tablosu"""
        elements = []
        elements.append(Paragraph("<b>4. NODE TABLOSU</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        # Tablo başlıkları
        data = [['Node ID', 'Tip', 'X (m)', 'Y (m)', 'Z (m)', 'Basınç (Bar)']]
        
        # Node verileri
        for node in self.network.nodes:
            row = [
                node.id,
                node.node_type,
                f"{node.x:.2f}",
                f"{node.y:.2f}",
                f"{node.elevation:.2f}",
                f"{getattr(node, 'pressure', 0):.2f}"
            ]
            data.append(row)
        
        t = Table(data, colWidths=[3*cm, 3*cm, 2.5*cm, 2.5*cm, 2.5*cm, 3*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(self.template.header_color)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        elements.append(t)
        
        return elements
    
    def _create_validation_report(self, styles):
        """Doğrulama raporu"""
        elements = []
        elements.append(Paragraph("<b>5. DOĞRULAMA RAPORU</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        if self.validation_results:
            elements.append(Paragraph(str(self.validation_results), styles['Normal']))
        else:
            elements.append(Paragraph("Doğrulama yapılmamış.", styles['Normal']))
        
        return elements
    
    def _create_material_list(self, styles):
        """Malzeme listesi"""
        elements = []
        elements.append(Paragraph("<b>6. MALZEME LİSTESİ</b>", styles['CustomHeading']))
        elements.append(Spacer(1, 0.5*cm))
        
        # Basit malzeme listesi - boruları grupla
        pipe_groups = {}
        for pipe in self.network.pipes:
            key = f"DN{pipe.diameter_mm:.0f}"
            if key not in pipe_groups:
                pipe_groups[key] = {'count': 0, 'total_length': 0}
            pipe_groups[key]['count'] += 1
            pipe_groups[key]['total_length'] += pipe.length
        
        data = [['Malzeme', 'Miktar', 'Birim']]
        
        for diameter, info in sorted(pipe_groups.items()):
            data.append([
                f"Boru {diameter}",
                f"{info['total_length']:.2f}",
                "m"
            ])
        
        # Sprinkler sayısı
        sprinkler_count = len([n for n in self.network.nodes if n.node_type == 'sprinkler'])
        if sprinkler_count > 0:
            data.append(['Sprinkler', str(sprinkler_count), 'adet'])
        
        t = Table(data, colWidths=[10*cm, 4*cm, 3*cm])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(self.template.header_color)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(t)
        
        return elements


class DOCXReportGenerator:
    """DOCX rapor oluşturucu"""
    
    def __init__(self, template: ReportTemplate, network, results, validation_results=None):
        self.template = template
        self.network = network
        self.results = results
        self.validation_results = validation_results
        
        if not HAS_DOCX:
            raise ImportError("python-docx kütüphanesi gerekli. 'pip install python-docx' çalıştırın.")
    
    def generate(self, filepath: str) -> bool:
        """DOCX rapor oluştur"""
        try:
            doc = Document()
            
            # Kapak sayfası
            if self.template.include_cover_page:
                self._create_cover_page(doc)
                doc.add_page_break()
            
            # Özet
            if self.template.include_summary:
                self._create_summary(doc)
            
            # Boru tablosu
            if self.template.include_pipe_table:
                self._create_pipe_table(doc)
            
            # Kaydet
            doc.save(filepath)
            return True
            
        except Exception as e:
            print(f"DOCX oluşturma hatası: {e}")
            return False
    
    def _create_cover_page(self, doc):
        """Kapak sayfası"""
        # Başlık
        title = doc.add_heading(self.template.project_name, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Alt başlık
        subtitle = doc.add_paragraph('Yangın Söndürme Sistemi Hidrolik Hesap Raporu')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        # Proje bilgileri
        if self.template.project_number:
            doc.add_paragraph(f'Proje No: {self.template.project_number}')
        if self.template.client_name:
            doc.add_paragraph(f'Müşteri: {self.template.client_name}')
        doc.add_paragraph(f'Tarih: {datetime.now().strftime("%d.%m.%Y")}')
    
    def _create_summary(self, doc):
        """Özet bölümü"""
        doc.add_heading('Özet', 1)
        
        doc.add_paragraph(f'Toplam Debi: {self.results.total_flow:.1f} L/dk')
        doc.add_paragraph(f'Toplam Basınç: {self.results.total_pressure:.2f} Bar')
        doc.add_paragraph(f'Maksimum Hız: {self.results.max_velocity:.2f} m/s')
    
    def _create_pipe_table(self, doc):
        """Boru tablosu"""
        doc.add_heading('Boru Tablosu', 1)
        
        # Tablo oluştur
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Light Grid Accent 1'
        
        # Başlıklar
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Boru ID'
        hdr_cells[1].text = 'Başlangıç'
        hdr_cells[2].text = 'Bitiş'
        hdr_cells[3].text = 'Çap (mm)'
        hdr_cells[4].text = 'Uzunluk (m)'
        hdr_cells[5].text = 'Debi (L/dk)'
        hdr_cells[6].text = 'Hız (m/s)'
        
        # Veri satırları
        for pipe in self.network.pipes:
            row_cells = table.add_row().cells
            row_cells[0].text = pipe.id
            row_cells[1].text = pipe.start_node_id
            row_cells[2].text = pipe.end_node_id
            row_cells[3].text = f"{pipe.diameter_mm:.0f}"
            row_cells[4].text = f"{pipe.length:.2f}"
            row_cells[5].text = f"{getattr(pipe, 'flow', 0):.1f}"
            row_cells[6].text = f"{getattr(pipe, 'velocity', 0):.2f}"


class ReportCustomizationDialog(tk.Toplevel):
    """Rapor özelleştirme dialogu"""
    
    def __init__(self, parent, network, results, validation_results=None):
        super().__init__(parent)
        self.network = network
        self.results = results
        self.validation_results = validation_results
        self.result = None
        
        self.title("Rapor Özelleştirme")
        self.geometry("800x700")
        
        # Modal
        self.transient(parent)
        self.grab_set()
        
        # Varsayılan template
        self.template = ReportTemplate(
            name="Standart Rapor",
            description="Tam kapsamlı standart rapor",
            format="pdf"
        )
        
        self._create_widgets()
        
        # Merkeze al
        self.update_idletasks()
        x = (self.winfo_screenwidth() // 2) - (self.winfo_width() // 2)
        y = (self.winfo_screenheight() // 2) - (self.winfo_height() // 2)
        self.geometry(f"+{x}+{y}")
    
    def _create_widgets(self):
        """Widget'ları oluştur"""
        # Ana frame
        main_frame = ttk.Frame(self, padding=10)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Notebook
        notebook = ttk.Notebook(main_frame)
        notebook.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        
        # Sekmeler
        self._create_general_tab(notebook)
        self._create_content_tab(notebook)
        self._create_company_tab(notebook)
        self._create_style_tab(notebook)
        
        # Butonlar
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X)
        
        ttk.Button(button_frame, text="PDF Oluştur", 
                  command=self._generate_pdf).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="DOCX Oluştur", 
                  command=self._generate_docx).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="İptal", 
                  command=self.destroy).pack(side=tk.RIGHT, padx=5)
    
    def _create_general_tab(self, notebook):
        """Genel ayarlar sekmesi"""
        frame = ttk.Frame(notebook, padding=10)
        notebook.add(frame, text="Genel")
        
        # Proje bilgileri
        ttk.Label(frame, text="Proje Adı:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.project_name_var = tk.StringVar(value=self.network.project_name)
        ttk.Entry(frame, textvariable=self.project_name_var, width=40).grid(row=0, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Proje No:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.project_number_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.project_number_var, width=40).grid(row=1, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Lokasyon:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.location_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.location_var, width=40).grid(row=2, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Müşteri:").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.client_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.client_var, width=40).grid(row=3, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Hazırlayan:").grid(row=4, column=0, sticky=tk.W, pady=5)
        self.engineer_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.engineer_var, width=40).grid(row=4, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Kontrol Eden:").grid(row=5, column=0, sticky=tk.W, pady=5)
        self.reviewer_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.reviewer_var, width=40).grid(row=5, column=1, sticky=tk.W)
    
    def _create_content_tab(self, notebook):
        """İçerik sekmesi"""
        frame = ttk.Frame(notebook, padding=10)
        notebook.add(frame, text="İçerik")
        
        # İçerik bölümleri checkbox'ları
        self.include_cover_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Kapak Sayfası", variable=self.include_cover_var).pack(anchor=tk.W, pady=2)
        
        self.include_toc_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="İçindekiler", variable=self.include_toc_var).pack(anchor=tk.W, pady=2)
        
        self.include_summary_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Özet", variable=self.include_summary_var).pack(anchor=tk.W, pady=2)
        
        self.include_calc_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Hesaplama Detayları", variable=self.include_calc_var).pack(anchor=tk.W, pady=2)
        
        self.include_pipes_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Boru Tablosu", variable=self.include_pipes_var).pack(anchor=tk.W, pady=2)
        
        self.include_nodes_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Node Tablosu", variable=self.include_nodes_var).pack(anchor=tk.W, pady=2)
        
        self.include_validation_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Doğrulama Raporu", variable=self.include_validation_var).pack(anchor=tk.W, pady=2)
        
        self.include_materials_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(frame, text="Malzeme Listesi", variable=self.include_materials_var).pack(anchor=tk.W, pady=2)
    
    def _create_company_tab(self, notebook):
        """Şirket bilgileri sekmesi"""
        frame = ttk.Frame(notebook, padding=10)
        notebook.add(frame, text="Şirket")
        
        ttk.Label(frame, text="Şirket Adı:").grid(row=0, column=0, sticky=tk.W, pady=5)
        self.company_name_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.company_name_var, width=50).grid(row=0, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Adres:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.company_address_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.company_address_var, width=50).grid(row=1, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Telefon:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.company_phone_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.company_phone_var, width=50).grid(row=2, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="E-posta:").grid(row=3, column=0, sticky=tk.W, pady=5)
        self.company_email_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.company_email_var, width=50).grid(row=3, column=1, sticky=tk.W)
        
        ttk.Label(frame, text="Web:").grid(row=4, column=0, sticky=tk.W, pady=5)
        self.company_web_var = tk.StringVar()
        ttk.Entry(frame, textvariable=self.company_web_var, width=50).grid(row=4, column=1, sticky=tk.W)
        
        # Logo seçimi
        ttk.Label(frame, text="Logo:").grid(row=5, column=0, sticky=tk.W, pady=5)
        logo_frame = ttk.Frame(frame)
        logo_frame.grid(row=5, column=1, sticky=tk.W)
        
        self.logo_path_var = tk.StringVar()
        ttk.Entry(logo_frame, textvariable=self.logo_path_var, width=35).pack(side=tk.LEFT)
        ttk.Button(logo_frame, text="Gözat...", command=self._select_logo).pack(side=tk.LEFT, padx=5)
    
    def _create_style_tab(self, notebook):
        """Stil ayarları sekmesi"""
        frame = ttk.Frame(notebook, padding=10)
        notebook.add(frame, text="Stil")
        
        # Renk seçimi
        ttk.Label(frame, text="Başlık Rengi:").grid(row=0, column=0, sticky=tk.W, pady=5)
        color_frame = ttk.Frame(frame)
        color_frame.grid(row=0, column=1, sticky=tk.W)
        
        self.header_color_var = tk.StringVar(value="#1f4788")
        color_entry = ttk.Entry(color_frame, textvariable=self.header_color_var, width=15)
        color_entry.pack(side=tk.LEFT)
        ttk.Button(color_frame, text="Seç...", 
                  command=lambda: self._select_color(self.header_color_var)).pack(side=tk.LEFT, padx=5)
        
        # Font boyutu
        ttk.Label(frame, text="Font Boyutu:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.font_size_var = tk.IntVar(value=10)
        ttk.Spinbox(frame, from_=8, to=14, textvariable=self.font_size_var, width=10).grid(row=1, column=1, sticky=tk.W)
    
    def _select_logo(self):
        """Logo dosyası seç"""
        filepath = filedialog.askopenfilename(
            title="Logo Seç",
            filetypes=[
                ("Image Files", "*.png *.jpg *.jpeg *.gif"),
                ("All Files", "*.*")
            ]
        )
        if filepath:
            self.logo_path_var.set(filepath)
    
    def _select_color(self, var):
        """Renk seç"""
        color = colorchooser.askcolor(title="Renk Seç", initialcolor=var.get())
        if color[1]:
            var.set(color[1])
    
    def _get_template(self) -> ReportTemplate:
        """Template oluştur"""
        template = ReportTemplate(
            name="Özel Rapor",
            description="Kullanıcı tarafından özelleştirilmiş rapor",
            format="pdf"
        )
        
        # Genel bilgiler
        template.project_name = self.project_name_var.get()
        template.project_number = self.project_number_var.get()
        template.project_location = self.location_var.get()
        template.client_name = self.client_var.get()
        template.engineer_name = self.engineer_var.get()
        template.reviewer_name = self.reviewer_var.get()
        
        # Şirket bilgileri
        template.company_name = self.company_name_var.get()
        template.company_address = self.company_address_var.get()
        template.company_phone = self.company_phone_var.get()
        template.company_email = self.company_email_var.get()
        template.company_web = self.company_web_var.get()
        template.logo_path = self.logo_path_var.get()
        
        # İçerik
        template.include_cover_page = self.include_cover_var.get()
        template.include_toc = self.include_toc_var.get()
        template.include_summary = self.include_summary_var.get()
        template.include_calculation_details = self.include_calc_var.get()
        template.include_pipe_table = self.include_pipes_var.get()
        template.include_node_table = self.include_nodes_var.get()
        template.include_validation_report = self.include_validation_var.get()
        template.include_material_list = self.include_materials_var.get()
        
        # Stil
        template.header_color = self.header_color_var.get()
        template.font_size = self.font_size_var.get()
        
        return template
    
    def _generate_pdf(self):
        """PDF rapor oluştur"""
        if not HAS_REPORTLAB:
            messagebox.showerror("Hata", "reportlab kütüphanesi gerekli!\n\npip install reportlab")
            return
        
        filepath = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF Files", "*.pdf"), ("All Files", "*.*")]
        )
        
        if not filepath:
            return
        
        template = self._get_template()
        generator = PDFReportGenerator(template, self.network, self.results, self.validation_results)
        
        try:
            success = generator.generate(filepath)
            if success:
                messagebox.showinfo("Başarılı", f"PDF rapor oluşturuldu:\n{filepath}")
                self.result = filepath
            else:
                messagebox.showerror("Hata", "PDF oluşturulamadı!")
        except Exception as e:
            messagebox.showerror("Hata", f"PDF oluşturma hatası:\n{str(e)}")
    
    def _generate_docx(self):
        """DOCX rapor oluştur"""
        if not HAS_DOCX:
            messagebox.showerror("Hata", "python-docx kütüphanesi gerekli!\n\npip install python-docx")
            return
        
        filepath = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")]
        )
        
        if not filepath:
            return
        
        template = self._get_template()
        generator = DOCXReportGenerator(template, self.network, self.results, self.validation_results)
        
        try:
            success = generator.generate(filepath)
            if success:
                messagebox.showinfo("Başarılı", f"DOCX rapor oluşturuldu:\n{filepath}")
                self.result = filepath
            else:
                messagebox.showerror("Hata", "DOCX oluşturulamadı!")
        except Exception as e:
            messagebox.showerror("Hata", f"DOCX oluşturma hatası:\n{str(e)}")


def show_report_customization(parent, network, results, validation_results=None):
    """Rapor özelleştirme dialogunu göster"""
    dialog = ReportCustomizationDialog(parent, network, results, validation_results)
    parent.wait_window(dialog)
    return dialog.result
